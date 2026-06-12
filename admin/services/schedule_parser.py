import re
import time
from datetime import datetime

from docx import Document
from flask import current_app

from models import db, Schedule

from .normalization import normalize_text, extract_time_pair_from_text
from .group_helpers import (
    parse_cell_multi_entries,
    get_or_create_group,
    build_group_lookup,
    extract_group_from_text,
)
from .teacher_helpers import (
    get_or_create_placeholder_teacher,
    build_teacher_lookup,
    find_teacher_in_text,
)
from .ai_parser import parse_multiple_lessons_with_ai
from .ai_client import AI_CHUNK_SIZE, AI_CALL_DELAY, AI_ENABLED


SUBJECT_ALIASES = {
    "КАЗ.ЯЗ": "Казахский язык",
    "КАЗ ЯЗ": "Казахский язык",
    "ҚАЗАҚ ТІЛІ": "Казахский язык",
    "КАЗАХСКИЙ ЯЗЫК": "Казахский язык",
    "РУС.ЯЗ": "Русский язык",
    "РУС ЯЗ": "Русский язык",
    "РУССКИЙ ЯЗЫК": "Русский язык",
    "РУС. ЯЗ": "Русский язык",
    "АНГ": "Английский язык",
    "АН.ЯЗ": "Английский язык",
    "АНГЛ": "Английский язык",
    "МАТЕМ": "Математика",
    "ФИЗ": "Физика",
    "ХИМ": "Химия",
    "БИО": "Биология",
    "ГЕО": "География",
    "ИНФ": "Информатика",
    "ИНФОРМ": "Информатика",
    "ИСТОРИЯ": "История",
    "ЛИТ": "Литература",
    "ФАКУЛЬТАТИВ": "Факультатив",
    "КОНСУЛЬТАЦИЯ": "Консультация",
}


def detect_subject(text: str) -> str:
    raw = normalize_text(text)
    if not raw:
        return "Предмет"

    normalized = raw.upper().replace("Ё", "Е")
    for alias, subject in SUBJECT_ALIASES.items():
        if alias.upper().replace("Ё", "Е") in normalized:
            return subject

    words = [w for w in raw.split() if len(w) > 2]
    if words:
        return " ".join(words[:4]).strip()

    return "Предмет"


def chunked(items, size):
    size = max(1, int(size))
    for i in range(0, len(items), size):
        yield items[i:i + size]


def fill_ai_nulls(ai_entry: dict, teachers_lookup: dict) -> dict:
    """
    Дозаполняет null из результата AI через локальные правила.
    """
    raw_text = ai_entry.get("raw", "") or ""

    group_raw = (ai_entry.get("group_raw") or "").strip()
    teacher_raw = (ai_entry.get("teacher_raw") or "").strip()
    subject_raw = (ai_entry.get("subject_raw") or "").strip()
    time_start = ai_entry.get("time_start")
    time_end = ai_entry.get("time_end")

    if not group_raw:
        group_raw, _ = extract_group_from_text(raw_text)

    if not teacher_raw:
        teacher_obj, _ = find_teacher_in_text(raw_text, teachers_lookup)
        if teacher_obj:
            teacher_raw = teacher_obj.name

    if not subject_raw:
        subject_raw = detect_subject(raw_text)

    return {
        "group_raw": group_raw,
        "teacher_raw": teacher_raw,
        "subject_raw": subject_raw,
        "time_start": time_start,
        "time_end": time_end,
        "raw": raw_text,
    }


def parse_and_save_schedule(filepath):
    doc = Document(filepath)

    stats = {
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "errors": 0,
        "fallback_teacher": 0,
        "ai_used": 0,
    }

    preview_items = []
    current_weekday = None
    date_pattern = r"(\d{1,2}\.\d{1,2}\.\d{4})"

    # Определение дня недели
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue

        m = re.search(date_pattern, text)
        if not m:
            continue

        try:
            date_obj = datetime.strptime(m.group(1), "%d.%m.%Y")
            current_weekday = date_obj.weekday() + 1
            current_app.logger.info(
                "Detected date: %s => weekday %s",
                m.group(1),
                current_weekday,
            )
            break
        except Exception:
            continue

    if not current_weekday:
        current_weekday = 1

    groups_lookup = build_group_lookup()
    teachers_lookup = build_teacher_lookup()
    placeholder_teacher = get_or_create_placeholder_teacher()

    for table in doc.tables:
        if len(table.rows) < 2:
            continue

        header = table.rows[0]
        column_times = {}

        # Время колонок
        for i, cell in enumerate(header.cells[1:], start=0):
            times = extract_time_pair_from_text(normalize_text(cell.text))
            if times:
                column_times[i] = times

        # Строки
        for row in table.rows[1:]:
            cells = row.cells
            if len(cells) < 2:
                continue

            room_name = normalize_text(cells[0].text).strip()
            if not room_name or room_name.upper() in ("АУД", "AUD"):
                continue

            room_name = re.sub(
                r"^([A-ZА-ЯЁ]{1,3})(\d+)",
                r"\1 \2",
                room_name,
                flags=re.IGNORECASE,
            )
            room_name = re.sub(r"\s+", " ", room_name).strip()

            pending_ai_entries = []

            for col_idx, cell in enumerate(cells[1:]):
                cell_text = normalize_text(cell.text)
                if not cell_text:
                    continue

                times = column_times.get(col_idx) or extract_time_pair_from_text(cell_text)
                if not times:
                    stats["skipped"] += 1
                    continue

                time_start, time_end = times

                # Быстрый regex-разбор
                entries = parse_cell_multi_entries(cell_text)

                if entries:
                    for entry in entries:
                        group_raw = entry.get("group_raw")
                        raw = entry.get("raw", "")

                        if not group_raw:
                            stats["skipped"] += 1
                            continue

                        group_obj = get_or_create_group(group_raw, groups_lookup)
                        if not group_obj:
                            stats["skipped"] += 1
                            continue

                        teacher_obj, _ = find_teacher_in_text(raw, teachers_lookup)
                        if not teacher_obj:
                            teacher_obj = placeholder_teacher
                            if teacher_obj:
                                stats["fallback_teacher"] += 1
                            else:
                                stats["skipped"] += 1
                                continue

                        subject_name = detect_subject(raw)

                        try:
                            with db.session.no_autoflush:
                                schedule_entry = Schedule.query.filter_by(
                                    group_id=group_obj.id,
                                    weekday=current_weekday,
                                    time_start=time_start,
                                ).first()

                            if schedule_entry:
                                schedule_entry.subject = subject_name
                                schedule_entry.teacher_id = teacher_obj.id
                                schedule_entry.room = room_name
                                schedule_entry.time_end = time_end
                                stats["updated"] += 1
                            else:
                                db.session.add(
                                    Schedule(
                                        group_id=group_obj.id,
                                        subject=subject_name,
                                        teacher_id=teacher_obj.id,
                                        room=room_name,
                                        weekday=current_weekday,
                                        time_start=time_start,
                                        time_end=time_end,
                                    )
                                )
                                stats["created"] += 1

                            preview_items.append({
                                "group": group_obj.name,
                                "teacher": teacher_obj.name,
                                "subject": subject_name,
                                "room": room_name,
                                "weekday": current_weekday,
                                "time_start": time_start,
                                "time_end": time_end,
                            })

                        except Exception as e:
                            db.session.rollback()
                            stats["errors"] += 1
                            current_app.logger.exception("Ошибка сохранения: %s", e)

                else:
                    pending_ai_entries.append({
                        "room": room_name,
                        "time_start": time_start,
                        "time_end": time_end,
                        "raw": cell_text,
                    })

            # Промежуточный commit после быстрого разбора строки
            try:
                db.session.commit()
            except Exception as e:
                db.session.rollback()
                current_app.logger.exception("Ошибка промежуточного коммита: %s", e)

            # AI только для тех ячеек, где regex не помог
            if not AI_ENABLED or not pending_ai_entries:
                continue

            for chunk in chunked(pending_ai_entries, AI_CHUNK_SIZE):
                ai_results = parse_multiple_lessons_with_ai(chunk)

                if ai_results:
                    stats["ai_used"] += 1

                for ai_entry in ai_results:
                    if not isinstance(ai_entry, dict):
                        continue

                    ai_entry = fill_ai_nulls(ai_entry, teachers_lookup)

                    group_raw = (ai_entry.get("group_raw") or "").strip()
                    teacher_raw = (ai_entry.get("teacher_raw") or "").strip()
                    subject_name = (ai_entry.get("subject_raw") or "").strip()
                    time_start = ai_entry.get("time_start")
                    time_end = ai_entry.get("time_end")
                    original_raw = ai_entry.get("raw", "")

                    if not group_raw or not time_start or not time_end:
                        stats["skipped"] += 1
                        continue

                    group_obj = get_or_create_group(group_raw, groups_lookup)
                    if not group_obj:
                        stats["skipped"] += 1
                        continue

                    teacher_obj = None
                    if teacher_raw:
                        teacher_obj, _ = find_teacher_in_text(teacher_raw, teachers_lookup)

                    if not teacher_obj:
                        teacher_obj, _ = find_teacher_in_text(original_raw, teachers_lookup)

                    if not teacher_obj:
                        teacher_obj = placeholder_teacher
                        if teacher_obj:
                            stats["fallback_teacher"] += 1
                        else:
                            stats["skipped"] += 1
                            continue

                    if not subject_name:
                        subject_name = detect_subject(original_raw)

                    try:
                        with db.session.no_autoflush:
                            exists = Schedule.query.filter_by(
                                group_id=group_obj.id,
                                weekday=current_weekday,
                                time_start=time_start,
                            ).first()

                        room_value = ai_entry.get("room") or room_name

                        if exists:
                            exists.subject = subject_name
                            exists.teacher_id = teacher_obj.id
                            exists.room = room_value
                            exists.time_end = time_end
                            stats["updated"] += 1
                        else:
                            db.session.add(
                                Schedule(
                                    group_id=group_obj.id,
                                    subject=subject_name,
                                    teacher_id=teacher_obj.id,
                                    room=room_value,
                                    weekday=current_weekday,
                                    time_start=time_start,
                                    time_end=time_end,
                                )
                            )
                            stats["created"] += 1

                        preview_items.append({
                            "group": group_obj.name,
                            "teacher": teacher_obj.name,
                            "subject": subject_name,
                            "room": room_value,
                            "weekday": current_weekday,
                            "time_start": time_start,
                            "time_end": time_end,
                        })

                    except Exception as e:
                        db.session.rollback()
                        stats["errors"] += 1
                        current_app.logger.exception("Ошибка AI сохранения: %s", e)

                if AI_CALL_DELAY > 0:
                    time.sleep(AI_CALL_DELAY)

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        current_app.logger.exception("Финальный commit failed: %s", e)

    return stats, preview_items